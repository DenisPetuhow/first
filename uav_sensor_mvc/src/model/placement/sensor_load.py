# -*- coding: utf-8 -*-
"""
MODEL · Нагрузка на средство (план 13.4) — ПОКАЗАТЕЛЬ, а не критерий: расстановка её не видит.

Нагрузка средства j — доля маршрутов выборки, прошедших через его зону: L_j = M_j / M × 100 %. «Прошёл» — хоть одна
точка маршрута не дальше радиуса средства, как засечка «Анализа» (`ThreatModel._detection`). Маршруты весят поровну,
средства равны (малые, у района защиты, большие). Разбор — `теория/код/НАГРУЗКА_НА_СРЕДСТВО.md`,
теория — `теория/пособие/5_КРИТЕРИИ_И_ПОКАЗАТЕЛИ.md`.

⚠️ ТОЛЬКО ЧИТАЕТ: расстановку, выборку и показатели не трогает — эталон от этого модуля не зависит.
"""
import numpy as np

# порядок строк таблицы (заказчик 30.09.2026: «настраиваемая сортировка, сначала базовая по убыванию нагрузки»):
# ключ · подпись в окне; первый — по умолчанию
ORDERS = (
    ("load_small_first", "нагрузка ↓, большие в конце"),
    ("load_desc", "нагрузка ↓, все вместе"),
    ("load_asc", "нагрузка ↑"),
    ("number", "№ по списку"),
    ("routes_desc", "маршрутов ↓"),
)


def seen_matrix(sample, pts, radii):
    """Какие маршруты прошли через какие средства — одна формула на модель: её зовут `_detection`, `sector_stats` и
    нагрузка. Вход: маршруты [(K,2) км], средства (N,2) км, радиусы (N,) км. Отдаёт: (M, N) bool."""
    pts = np.asarray(pts, float).reshape(-1, 2)      # средства, км
    rad = np.asarray(radii, float).reshape(-1)       # их радиусы, км
    out = np.zeros((len(sample), len(pts)), bool)    # прошёл ли маршрут через средство
    if not len(pts):                                 # средств нет
        return out                                   # — матрица без столбцов
    for i, r in enumerate(sample):                   # маршрут — строка матрицы
        xy = np.asarray(r, float).reshape(-1, 2)     # точки маршрута, км
        if len(xy):                                  # пустой маршрут никого не проходит
            out[i] = np.linalg.norm(xy[:, None, :] - pts[None, :, :], axis=2).min(axis=0) <= rad
    return out


def load_rows(table, seen, radii, big_id=4):
    """Нагрузка каждого средства. Вход: строки `sensors_table()` в порядке столбцов `seen` (M, N), радиусы (N,) км,
    номер большого типа. Отдаёт: dict(n, type_id, big, routes, load_pct, x_km, y_km, r_km) в порядке таблицы средств."""
    m = int(seen.shape[0])                           # маршрутов в выборке
    cnt = seen.sum(axis=0)                           # маршрутов через каждое средство
    rad = np.asarray(radii, float).reshape(-1)       # радиусы средств, км — для подсветки на карте
    return [dict(n=int(t["n"]), type_id=int(t["type_id"]), big=int(t["type_id"]) == int(big_id),
                 routes=int(cnt[j]), load_pct=(100.0 * float(cnt[j]) / m) if m else 0.0,
                 x_km=float(t["x_km"]), y_km=float(t["y_km"]), r_km=float(rad[j]))
            for j, t in enumerate(table)]


def sort_rows(rows, order=ORDERS[0][0]):
    """Строки в заданном порядке. Вход: строки `load_rows`, ключ из `ORDERS`. Отдаёт: новый список;
    при равенстве — по номеру, поэтому порядок повторяется до бита. Чужой ключ — порядок по умолчанию."""
    keys = {                                         # ключ порядка -> ключ сортировки строки
        "load_small_first": lambda r: (r["big"], -r["load_pct"], r["n"]),   # малые, затем большие
        "load_desc": lambda r: (-r["load_pct"], r["n"]),
        "load_asc": lambda r: (r["load_pct"], r["n"]),
        "number": lambda r: r["n"],
        "routes_desc": lambda r: (-r["routes"], r["n"]),
    }
    return sorted(rows, key=keys.get(order, keys[ORDERS[0][0]]))


def write_docx(path, title, info, headers, cells):
    """Таблица нагрузки в Word. Вход: путь .docx, заголовок, строки пояснения, подписи столбцов, ячейки [[str]].
    Отдаёт: None — записано; иначе текст ошибки (нет python-docx, файл занят другой программой)."""
    try:
        from docx import Document
    except ImportError:                              # библиотеки нет — сказать, а не упасть
        return "нет библиотеки python-docx (pip install python-docx)"
    doc = Document()                                 # новый документ
    doc.add_heading(title, level=1)
    for line in info:                                # пояснение: выборка, тип, порядок, дата
        doc.add_paragraph(line)
    tab = doc.add_table(rows=1, cols=len(headers))   # таблица: строка подписей и строки средств
    tab.style = "Table Grid"
    for c, h in enumerate(headers):                  # подписи столбцов — жирным
        cell = tab.rows[0].cells[c]                  # ячейка подписи
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for row in cells:                                # строка таблицы на средство
        tr = tab.add_row().cells                     # ячейки новой строки
        for c, v in enumerate(row):                  # значения по столбцам
            tr[c].text = str(v)
    try:
        doc.save(path)
    except OSError as e:                             # файл открыт в Word или нет прав
        return "не записан: %s" % e
    return None
